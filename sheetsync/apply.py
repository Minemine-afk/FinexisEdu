"""Read the editors' workbook and write their changes back into data/.

    python -m sheetsync.apply finexisedu-fee-sheet.xlsx            # write changes
    python -m sheetsync.apply finexisedu-fee-sheet.xlsx --dry-run   # only report

All or nothing: if any row has a problem, no file is touched and the report lists every problem
with its sheet and row so the editor can fix it. Only cells that were changed are checked against
the writing rules, so existing text is never rejected just for being old.
"""
import argparse
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from . import build, jsonio
from . import layout as L

CURRENCY_MARK = re.compile(r"[£$€¥]|\b[A-Z]{3}\s?\d|\d{3,}")


@dataclass
class Result:
    changes: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    written: list[str] = field(default_factory=list)


class Sheet:
    """One worksheet read as rows keyed by header name."""

    def __init__(self, ws, name: str, result: Result):
        self.name = name
        header = [c.value for c in ws[1]]
        self.rows: list[tuple[int, dict]] = []
        expected = [L.KEY, *L.IDENTITY[name], *L.EDITABLE[name]]
        missing = [h for h in expected if h not in header]
        if missing:
            result.problems.append(f"{name}: column(s) missing or renamed: {', '.join(missing)}")
            return
        for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            self.rows.append((n, dict(zip(header, row))))


class Applier:
    def __init__(self, root: Path):
        self.root = root
        self.result = Result()
        self.files: dict[str, tuple[Path, dict, str]] = {}
        self.dirty: set[str] = set()
        for path in sorted((root / "data" / "universities").glob("*/*.json")):
            data, text = jsonio.read(path)
            self.files[f"{path.parent.name}/{path.stem}"] = (path, data, text)

    # ---- helpers -------------------------------------------------------
    def problem(self, sheet: str, row: int, msg: str) -> None:
        self.result.problems.append(f"{sheet}, row {row}: {msg}")

    def change(self, sheet: str, row: int, label: str, what: str, old, new) -> None:
        self.result.changes.append(f"{sheet}, row {row}: {label}: {what} {self.show(old)} -> {self.show(new)}")

    @staticmethod
    def show(v) -> str:
        s = "(empty)" if v is None else str(v)
        return s if len(s) <= 70 else s[:67] + "..."

    def number(self, sheet, row, label, v, *, required=True, integer=False):
        if v is None or (isinstance(v, str) and not v.strip()):
            if required:
                self.problem(sheet, row, f"{label} is empty")
            return None
        if isinstance(v, bool):
            self.problem(sheet, row, f"{label} must be a number")
            return None
        if isinstance(v, str):
            try:
                v = float(v.replace(",", "").strip())
            except ValueError:
                self.problem(sheet, row, f"{label} must be a number, not '{v}'")
                return None
        if v < 0:
            self.problem(sheet, row, f"{label} cannot be negative")
            return None
        v = round(float(v), 2)
        if integer:
            if v != int(v):
                self.problem(sheet, row, f"{label} must be a whole number")
                return None
            return int(v)
        return int(v) if v == int(v) else v

    def text(self, v):
        if v is None:
            return None
        s = str(v).replace("\r\n", "\n").strip()
        return s or None

    def day(self, sheet, row, label, v):
        if isinstance(v, (datetime, date)):
            return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
        s = self.text(v)
        if s is None:
            self.problem(sheet, row, f"{label} is empty")
            return None
        try:
            date.fromisoformat(s)
        except ValueError:
            self.problem(sheet, row, f"{label} must look like 2026-10-21, not '{s}'")
            return None
        return s

    def source_type(self, sheet, row, v, default="official"):
        s = self.text(v) or default
        if s not in L.SOURCE_TYPES:
            self.problem(sheet, row, f"Source type must be official or secondary, not '{s}'")
            return None
        return s

    def url(self, sheet, row, v, *, required):
        s = self.text(v)
        if s is None:
            if required:
                self.problem(sheet, row, "Source URL is empty")
            return None
        if not re.match(r"https?://\S+$", s):
            self.problem(sheet, row, "Source URL must start with http:// or https:// and have no spaces")
            return None
        return s

    def set_key(self, obj: dict, key: str, new, uid: str, *, sheet, row, label, what) -> None:
        """Write obj[key] = new (or remove it when new is None) and log it when it differs."""
        old = obj.get(key)
        if old == new:
            return
        if new is None:
            del obj[key]
        else:
            obj[key] = new
        self.dirty.add(uid)
        self.change(sheet, row, label, what, old, new)

    def locate(self, sheet, row, key, uid):
        if uid not in self.files:
            self.problem(sheet, row, f"unknown university in key '{key}'")
            return None
        return self.files[uid][1]

    def check_identity(self, sheet, row, values: dict, expected: dict) -> bool:
        for col, want in expected.items():
            got = values.get(col)
            if str(got).strip() != str(want):
                self.problem(sheet, row, f"{col} is '{got}' but this row's key belongs to '{want}'. "
                                         "Rows were probably sorted or moved; restore the sheet and edit again")
                return False
        return True

    def duplicates(self, sheet: Sheet) -> set:
        seen, keys = set(), set()
        for n, r in sheet.rows:
            k = r.get(L.KEY)
            if k in seen:
                self.problem(sheet.name, n, f"key '{k}' appears twice")
            seen.add(k)
            keys.add(k)
        return keys

    def expect_keys(self, sheet: str, have: set, want: set) -> None:
        for k in sorted(want - have):
            self.result.problems.append(f"{sheet}: the row for '{k}' was deleted; rows cannot be removed")

    # ---- fees ----------------------------------------------------------
    def apply_fees(self, sheet: Sheet) -> None:
        have = self.duplicates(sheet)
        unis = build.load_universities(self.root)
        self.expect_keys(sheet.name, have, {r[0] for r in build.fee_rows(unis)})
        prog_edits: dict[tuple, dict[str, list]] = {}
        for n, r in sheet.rows:
            key = str(r[L.KEY])
            parts = key.split("|")
            if len(parts) != 5:
                self.problem(sheet.name, n, f"unrecognised key '{key}'")
                continue
            uid, pi, kind, tier, year = parts
            data = self.locate(sheet.name, n, key, uid)
            if data is None:
                continue
            try:
                p = data["programmes"][int(pi)]
            except (ValueError, IndexError):
                self.problem(sheet.name, n, f"unknown programme in key '{key}'")
                continue
            label = f"{data['name']} / {build.programme_label(p)} / {tier} / {year}"
            ident = {"University": data["name"], "Programme": build.programme_label(p), "Tier": tier,
                     "Row type": {"cur": "current", "hist": "earlier", "later": "later"}.get(kind),
                     "Fee year": year if kind != "later" else p["feeYear"]}
            if not self.check_identity(sheet.name, n, r, ident):
                continue
            tuition = self.number(sheet.name, n, "Annual tuition", r["Annual tuition"])
            if kind == "cur":
                fv = p["fees"].get(tier)
                if fv is None:
                    self.problem(sheet.name, n, f"no {tier} fee in the data for key '{key}'")
                    continue
                comp = self.number(sheet.name, n, "Compulsory fees", r["Compulsory fees"])
                one = self.number(sheet.name, n, "One-off fees", r["One-off fees"])
                if None in (tuition, comp, one):
                    continue
                for k, v in (("annualTuition", tuition), ("annualCompulsoryFees", comp), ("oneOffFees", one)):
                    self.set_key(fv, k, v, uid, sheet=sheet.name, row=n, label=label, what=k)
                edits = prog_edits.setdefault((uid, int(pi)), {})
                stype = self.source_type(sheet.name, n, r["Source type"], p.get("sourceType", "official"))
                url = self.url(sheet.name, n, r["Source URL"], required=True)
                checked = self.day(sheet.name, n, "Last verified", r["Last verified"])
                for k, v in (("sourceType", stype), ("sourceUrl", url), ("lastVerified", checked)):
                    if v is not None:
                        edits.setdefault(k, []).append((n, v))
            elif kind == "hist":
                entry = next((h for h in p.get("feeHistory", [])
                              if h["tier"] == tier and str(h["feeYear"]) == year), None)
                if entry is None:
                    self.problem(sheet.name, n, f"no earlier-year fee in the data for key '{key}'")
                    continue
                if tuition is None:
                    continue
                comp = self.number(sheet.name, n, "Compulsory fees", r["Compulsory fees"], required=False)
                one = self.number(sheet.name, n, "One-off fees", r["One-off fees"], required=False)
                stype = self.source_type(sheet.name, n, r["Source type"], entry.get("sourceType", "official"))
                url = self.url(sheet.name, n, r["Source URL"], required=False)
                if stype is None:
                    continue
                if stype == "official" and "sourceType" not in entry:
                    stype = None
                for k, v in (("annualTuition", tuition), ("annualCompulsoryFees", comp), ("oneOffFees", one),
                             ("sourceType", stype if stype is not None else entry.get("sourceType")),
                             ("sourceUrl", url)):
                    self.set_key(entry, k, v, uid, sheet=sheet.name, row=n, label=label, what=k)
                if self.text(r["Last verified"]) is not None:
                    self.result.warnings.append(f"{sheet.name}, row {n}: Last verified on an earlier-year row is ignored")
            else:
                entry = next((x for x in p.get("laterYears", [])
                              if x.get("tier", "all") == tier and str(x["fromYear"]) == year.split("|")[0]), None)
                if entry is None:
                    self.problem(sheet.name, n, f"no later-year fee in the data for key '{key}'")
                    continue
                if tuition is None:
                    continue
                comp = self.number(sheet.name, n, "Compulsory fees", r["Compulsory fees"], required=False)
                for k, v in (("annualTuition", tuition), ("annualCompulsoryFees", comp)):
                    self.set_key(entry, k, v, uid, sheet=sheet.name, row=n, label=label, what=k)
        for (uid, pi), fields in prog_edits.items():
            p = self.files[uid][1]["programmes"][pi]
            for k, found in fields.items():
                changed = {v for _, v in found if v != p.get(k)}
                if len(changed) > 1:
                    rows = ", ".join(str(n) for n, _ in found)
                    self.problem(sheet.name, found[0][0],
                                 f"rows {rows} are the same programme but give different {k} values; use one value")
                elif changed:
                    n = next(n for n, v in found if v in changed)
                    self.set_key(p, k, changed.pop(), uid, sheet=sheet.name, row=n,
                                 label=f"{self.files[uid][1]['name']} / {build.programme_label(p)}", what=k)

    # ---- living costs --------------------------------------------------
    def apply_living(self, sheet: Sheet) -> None:
        have = self.duplicates(sheet)
        self.expect_keys(sheet.name, have, {r[0] for r in build.living_rows(build.load_universities(self.root))})
        for n, r in sheet.rows:
            key = str(r[L.KEY])
            uid = key.split("|")[0]
            data = self.locate(sheet.name, n, key, uid)
            if data is None:
                continue
            k = data.get("livingCosts")
            if k is None or not key.endswith("|living"):
                self.problem(sheet.name, n, f"no living costs in the data for key '{key}'")
                continue
            if not self.check_identity(sheet.name, n, r, {"University": data["name"]}):
                continue
            label = f"{data['name']} living costs"
            year = self.number(sheet.name, n, "Year", r["Year"], integer=True)
            months = self.number(sheet.name, n, "Months a year", r["Months a year"])
            if months is not None and not 1 <= months <= 12:
                self.problem(sheet.name, n, "Months a year must be 1 to 12")
                months = None
            monthly = {name: self.number(sheet.name, n, col, r[col]) for name, col in (
                ("housing", "Housing / month"), ("food", "Food / month"),
                ("transport", "Transport / month"), ("personal", "Personal / month"))}
            stype = self.source_type(sheet.name, n, r["Source type"], k.get("sourceType", "official"))
            url = self.url(sheet.name, n, r["Source URL"], required=True)
            checked = self.day(sheet.name, n, "Last verified", r["Last verified"])
            if None in (year, months, stype, url, checked, *monthly.values()):
                continue
            for name, v in (("year", year), ("months", months), ("sourceType", stype), ("sourceUrl", url),
                            ("lastVerified", checked)):
                self.set_key(k, name, v, uid, sheet=sheet.name, row=n, label=label, what=name)
            for name, v in monthly.items():
                self.set_key(k["monthly"], name, v, uid, sheet=sheet.name, row=n, label=label, what=name)

    # ---- text ----------------------------------------------------------
    def check_text(self, sheet, row, label, old, new, *, limit, required, one_sentence=False, no_amounts=False):
        """Return the cleaned text, or the old value untouched after logging a problem."""
        if new == self.text(old):
            return old  # same text apart from stray spaces or line endings
        if new is None:
            if required:
                self.problem(sheet, row, f"{label} cannot be empty")
                return old
            return None
        if len(new) > limit:
            self.problem(sheet, row, f"{label} is {len(new)} characters; the limit is {limit}")
            return old
        if re.search(r"https?://|www\.", new) and (no_amounts or label.startswith("How")):
            self.problem(sheet, row, f"{label} must not contain a web address")
            return old
        if one_sentence and re.search(r"[.!?]\s+\S", new):
            self.problem(sheet, row, f"{label} must be one sentence")
            return old
        if no_amounts and CURRENCY_MARK.search(new):
            self.problem(sheet, row, f"{label} must not contain amounts or currency (those belong in the summary)")
            return old
        return new

    def apply_university_text(self, sheet: Sheet) -> None:
        have = self.duplicates(sheet)
        self.expect_keys(sheet.name, have, {r[0] for r in build.university_text_rows(build.load_universities(self.root))})
        fields = [("Description (shown on hover)", "description", 1500, "Description"),
                  ("Why included (max 400)", "whyIncluded", 400, "Why included"),
                  ("How fees are priced (max 220, not shown)", "feeMethod", 220, "How fees are priced")]
        for n, r in sheet.rows:
            key = str(r[L.KEY])
            uid = key.split("|")[0]
            data = self.locate(sheet.name, n, key, uid)
            if data is None or not self.check_identity(sheet.name, n, r, {"University": data["name"]}):
                continue
            for col, name, limit, label in fields:
                new = self.check_text(sheet.name, n, label, data[name], self.text(r[col]), limit=limit, required=True)
                self.set_key(data, name, new, uid, sheet=sheet.name, row=n, label=data["name"], what=name)

    def apply_programme_text(self, sheet: Sheet) -> None:
        have = self.duplicates(sheet)
        self.expect_keys(sheet.name, have, {r[0] for r in build.programme_text_rows(build.load_universities(self.root))})
        for n, r in sheet.rows:
            key = str(r[L.KEY])
            parts = key.split("|")
            data = self.locate(sheet.name, n, key, parts[0])
            if data is None:
                continue
            try:
                p = data["programmes"][int(parts[1])]
            except (ValueError, IndexError):
                self.problem(sheet.name, n, f"unknown programme in key '{key}'")
                continue
            if not self.check_identity(sheet.name, n, r, {"University": data["name"],
                                                          "Programme": build.programme_label(p)}):
                continue
            label = f"{data['name']} / {build.programme_label(p)}"
            for col, name, limit, req, opts in (
                ("Summary (shown, max 280)", "summary", 280, True, {}),
                ("How this fee was worked out (shown, max 160, one sentence, no amounts)", "feeChoice", 160, True,
                 {"one_sentence": True, "no_amounts": True}),
                ("Fee basis (not shown, max 240)", "feeBasis", 240, True, {}),
                ("Notes (not shown)", "notes", 100000, False, {}),
            ):
                new = self.check_text(sheet.name, n, {"summary": "Summary", "feeChoice": "How this fee was worked out",
                                                      "feeBasis": "Fee basis", "notes": "Notes"}[name],
                                      p.get(name), self.text(r[col]), limit=limit, required=req, **opts)
                self.set_key(p, name, new, parts[0], sheet=sheet.name, row=n, label=label, what=name)

    # ---- run -----------------------------------------------------------
    def run(self, workbook: Path, dry_run: bool) -> Result:
        wb = load_workbook(workbook, data_only=True)
        handlers = {L.FEES: self.apply_fees, L.LIVING: self.apply_living,
                    L.UNIVERSITY_TEXT: self.apply_university_text, L.PROGRAMME_TEXT: self.apply_programme_text}
        for name, handler in handlers.items():
            if name not in wb.sheetnames:
                self.result.problems.append(f"The sheet '{name}' is missing")
                continue
            sheet = Sheet(wb[name], name, self.result)
            if sheet.rows:
                handler(sheet)
        if self.result.problems or dry_run:
            return self.result
        for uid in sorted(self.dirty):
            path, data, text = self.files[uid]
            path.write_text(jsonio.dumps(data, text))
            self.result.written.append(str(path.relative_to(self.root)))
        return self.result


def report(result: Result, dry_run: bool) -> str:
    out = ["# Sheet sync", ""]
    if result.problems:
        out += [f"**Nothing was changed. {len(result.problems)} problem(s) to fix in the sheet:**", ""]
        out += [f"- {p}" for p in result.problems]
    else:
        verb = "would change" if dry_run else "changed"
        out += [f"{len(result.changes)} cell(s) {verb}" + (f" in {len(result.written)} file(s)." if result.written else "."), ""]
        out += [f"- {c}" for c in result.changes[:200]]
        if len(result.changes) > 200:
            out.append(f"- ... and {len(result.changes) - 200} more")
    if result.warnings:
        out += ["", "Warnings:", *[f"- {w}" for w in result.warnings]]
    return "\n".join(out) + "\n"


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--root", type=Path, default=Path("."))
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--report", type=Path, help="also write the report to this file")
    args = ap.parse_args(argv)
    result = Applier(args.root).run(args.workbook, args.dry_run)
    text = report(result, args.dry_run)
    print(text)
    if args.report:
        args.report.write_text(text)
    return 1 if result.problems else 0


if __name__ == "__main__":
    sys.exit(main())
