"""Write the editors' workbook from the data/ files.

    python -m sheetsync.build . finexisedu-fee-sheet.xlsx
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter as col_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import layout as L

LOCKED_HEAD = PatternFill("solid", fgColor="1D4D3F")
EDIT_HEAD = PatternFill("solid", fgColor="1D7A5A")
GREY = PatternFill("solid", fgColor="EEF2F0")

READ_ME = [
    "FinexisEdu fee sheet (edit existing figures only)",
    "",
    "What you can do: change the green-headed columns on the Fees, Living costs, University text and Programme text tabs.",
    "What you cannot do: add, delete or rename universities, programmes or rows. The grey columns are locked.",
    "",
    "Rules for fees:",
    "- If a university publishes a range, enter the TOP of the range (current and earlier years).",
    "- Text rules: Summary max 280 characters; the 'How this fee was worked out' line is one sentence, max 160, with no amounts, currency or URLs; Why included max 400 and should link the university to Singapore.",
    "- Only enter figures you found in a source. Never back-calculate. Put the page address in Source URL.",
    "- Source type: official for the university or government site, secondary only for a dated third-party page.",
    "- Update Last verified (YYYY-MM-DD) when you check a figure.",
    "- Leave Compulsory fees and One-off fees empty on earlier years to mean 'same as the current year'.",
    "",
    "Row types on the Fees tab:",
    "- current: the fee for the programme's current fee year.",
    "- earlier: a fee for an earlier year (history).",
    "- later: a different fee for later years of study (for example clinical years).",
    "",
    "Do not edit the Key column. Use the filter buttons rather than sorting, so every row keeps its own key.",
]


def load_universities(root: Path) -> list[dict]:
    unis = []
    for path in sorted((root / "data" / "universities").glob("*/*.json")):
        u = json.loads(path.read_text())
        u["_uid"] = f"{path.parent.name}/{path.stem}"
        unis.append(u)
    unis.sort(key=lambda u: (u["country"], u["name"]))
    return unis


def programme_label(p: dict) -> str:
    return p["name"] + (f" ({p['partner']})" if p.get("partner") else "")


def fee_rows(unis: list[dict]):
    for u in unis:
        for pi, p in enumerate(u["programmes"]):
            label = programme_label(p)
            for tier, fv in p["fees"].items():
                yield [f"{u['_uid']}|{pi}|cur|{tier}|{p['feeYear']}", u["name"], label, p["level"], tier, "current",
                       p["feeYear"], None, u["currency"], fv["annualTuition"], fv.get("annualCompulsoryFees", 0),
                       fv.get("oneOffFees", 0), p.get("sourceType", "official"), p.get("sourceUrl"), p.get("lastVerified")]
            for h in p.get("feeHistory", []):
                yield [f"{u['_uid']}|{pi}|hist|{h['tier']}|{h['feeYear']}", u["name"], label, p["level"], h["tier"],
                       "earlier", h["feeYear"], None, u["currency"], h["annualTuition"], h.get("annualCompulsoryFees"),
                       h.get("oneOffFees"), h.get("sourceType", "official"), h.get("sourceUrl"), None]
            for lt in p.get("laterYears", []):
                tier = lt.get("tier", "all")
                yield [f"{u['_uid']}|{pi}|later|{tier}|{lt['fromYear']}", u["name"], label, p["level"], tier, "later",
                       p["feeYear"], lt["fromYear"], u["currency"], lt["annualTuition"],
                       lt.get("annualCompulsoryFees"), None, None, None, None]


def living_rows(unis: list[dict]):
    for u in unis:
        k = u.get("livingCosts")
        if not k:
            continue
        m = k["monthly"]
        yield [f"{u['_uid']}|living", u["name"], u["country"], u["currency"], k["year"], k["months"], m["housing"],
               m["food"], m["transport"], m["personal"], k.get("sourceType", "official"), k.get("sourceUrl"),
               k.get("lastVerified")]


def university_text_rows(unis: list[dict]):
    for u in unis:
        yield [f"{u['_uid']}|uni", u["name"], u["country"], u["city"], u["description"], u["whyIncluded"],
               u["feeMethod"]]


def programme_text_rows(unis: list[dict]):
    for u in unis:
        for pi, p in enumerate(u["programmes"]):
            yield [f"{u['_uid']}|{pi}|prog", u["name"], programme_label(p), p["level"], p["summary"], p["feeChoice"],
                   p["feeBasis"], p.get("notes")]


FEES_COLS = [("Key", 22), ("University", 28), ("Programme", 40), ("Level", 10), ("Tier", 13), ("Row type", 10),
             ("Fee year", 9), ("From study year", 10), ("Currency", 9), ("Annual tuition", 14),
             ("Compulsory fees", 13), ("One-off fees", 12), ("Source type", 12), ("Source URL", 50),
             ("Last verified", 13)]
LIVING_COLS = [("Key", 22), ("University", 30), ("Country", 9), ("Currency", 9), ("Year", 8), ("Months a year", 9),
               ("Housing / month", 14), ("Food / month", 13), ("Transport / month", 14), ("Personal / month", 14),
               ("Source type", 12), ("Source URL", 50), ("Last verified", 13)]
UNI_TEXT_COLS = [("Key", 20), ("University", 30), ("Country", 9), ("City", 14), ("Description (shown on hover)", 60),
                 ("Why included (max 400)", 60), ("How fees are priced (max 220, not shown)", 50)]
PROG_TEXT_COLS = [("Key", 22), ("University", 26), ("Programme", 34), ("Level", 9), ("Summary (shown, max 280)", 55),
                  ("How this fee was worked out (shown, max 160, one sentence, no amounts)", 50),
                  ("Fee basis (not shown, max 240)", 50), ("Notes (not shown)", 50)]


def add_sheet(wb, title, cols, rows, validations, wrap=False, freeze="C2"):
    ws = wb.create_sheet(title)
    editable = set(L.EDITABLE[title])
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(1, i, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = EDIT_HEAD if name in editable else LOCKED_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[col_letter(i)].width = width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = freeze
    n = 1
    for row in rows:
        n += 1
        for i, v in enumerate(row, 1):
            c = ws.cell(n, i, v)
            if wrap:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if cols[i - 1][0] in editable:
                c.protection = Protection(locked=False)
            else:
                c.fill = GREY
            if cols[i - 1][0] == "Last verified":
                c.number_format = "@"
    index = {name: col_letter(i) for i, (name, _) in enumerate(cols, 1)}
    for names, kwargs in validations:
        dv = DataValidation(showErrorMessage=True, **kwargs)
        for name in names:
            dv.add(f"{index[name]}2:{index[name]}{n}")
        ws.add_data_validation(dv)
    ws.auto_filter.ref = f"A1:{col_letter(len(cols))}{n}"
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = True
    return n - 1


def number_rule(names):
    return (names, dict(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        errorTitle="Not allowed", error="Enter a number of 0 or more."))


def source_rule():
    return (["Source type"], dict(type="list", formula1='"official,secondary"', allow_blank=True,
                                  errorTitle="Not allowed", error="Choose official or secondary."))


def length_rule(name, limit):
    return ([name], dict(type="textLength", operator="lessThanOrEqual", formula1=str(limit), allow_blank=False,
                         errorTitle="Too long", error=f"Maximum {limit} characters."))


def build(root: Path, out: Path) -> dict[str, int]:
    unis = load_universities(Path(root))
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    for i, line in enumerate(READ_ME, 1):
        readme.cell(i, 1, line)
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 110
    counts = {
        L.FEES: add_sheet(wb, L.FEES, FEES_COLS, fee_rows(unis),
                          [number_rule(["Annual tuition", "Compulsory fees", "One-off fees"]), source_rule()],
                          freeze="F2"),
        L.LIVING: add_sheet(wb, L.LIVING, LIVING_COLS, living_rows(unis),
                            [number_rule(["Housing / month", "Food / month", "Transport / month", "Personal / month"]),
                             (["Months a year"], dict(type="decimal", operator="between", formula1="1", formula2="12",
                                                      errorTitle="Not allowed", error="Months a year must be 1 to 12.")),
                             source_rule()]),
        L.UNIVERSITY_TEXT: add_sheet(wb, L.UNIVERSITY_TEXT, UNI_TEXT_COLS, university_text_rows(unis),
                                     [length_rule("Why included (max 400)", 400),
                                      length_rule("How fees are priced (max 220, not shown)", 220)], wrap=True),
        L.PROGRAMME_TEXT: add_sheet(wb, L.PROGRAMME_TEXT, PROG_TEXT_COLS, programme_text_rows(unis),
                                    [length_rule("Summary (shown, max 280)", 280),
                                     length_rule("How this fee was worked out (shown, max 160, one sentence, no amounts)", 160),
                                     length_rule("Fee basis (not shown, max 240)", 240)], wrap=True),
    }
    wb.save(out)
    return counts


if __name__ == "__main__":
    print(build(Path(sys.argv[1]), Path(sys.argv[2])))
