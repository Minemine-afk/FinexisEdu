import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from sheetsync import build
from sheetsync import layout as L
from sheetsync.apply import Applier

REPO = Path(__file__).resolve().parents[2]
UNIVERSITIES = ["sg/ntu", "uk/manchester", "uk/lse"]


@pytest.fixture()
def root(tmp_path):
    for uid in UNIVERSITIES:
        dest = tmp_path / "data" / "universities" / f"{uid}.json"
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(REPO / "data" / "universities" / f"{uid}.json", dest)
    return tmp_path


@pytest.fixture()
def workbook(root, tmp_path):
    path = tmp_path / "sheet.xlsx"
    build.build(root, path)
    return path


def find(ws, key):
    for row in ws.iter_rows(min_row=2):
        if row[0].value == key:
            return row[0].row
    raise KeyError(key)


def column(ws, header):
    return next(c.column for c in ws[1] if c.value == header)


def edit(workbook, sheet, key, header, value):
    wb = load_workbook(workbook)
    ws = wb[sheet]
    ws.cell(find(ws, key), column(ws, header)).value = value
    wb.save(workbook)


def run(root, workbook, dry_run=False):
    return Applier(root).run(workbook, dry_run)


def snapshot(root):
    return {p: p.read_text() for p in sorted(root.rglob("*.json"))}


def test_untouched_sheet_changes_nothing(root, workbook):
    before = snapshot(root)
    result = run(root, workbook)
    assert result.problems == [] and result.changes == []
    assert snapshot(root) == before


def test_fee_and_source_edits_reach_the_data(root, workbook):
    key = "uk/manchester|0|cur|international|2027"
    edit(workbook, L.FEES, key, "Annual tuition", 40000)
    edit(workbook, L.FEES, key, "Last verified", "2026-10-21")
    result = run(root, workbook)
    assert result.problems == []
    import json
    prog = json.loads((root / "data/universities/uk/manchester.json").read_text())["programmes"][0]
    assert prog["fees"]["international"]["annualTuition"] == 40000
    assert prog["lastVerified"] == "2026-10-21"
    assert result.written == ["data/universities/uk/manchester.json"]


def test_earlier_year_blank_compulsory_fee_is_removed(root, workbook):
    import json
    key = "uk/manchester|0|hist|international|2025"
    edit(workbook, L.FEES, key, "Annual tuition", 36500)
    result = run(root, workbook)
    assert result.problems == []
    prog = json.loads((root / "data/universities/uk/manchester.json").read_text())["programmes"][0]
    entry = next(h for h in prog["feeHistory"] if h["feeYear"] == 2025)
    assert entry["annualTuition"] == 36500 and "annualCompulsoryFees" not in entry


def test_text_edits_and_untouched_layout(root, workbook):
    import json
    edit(workbook, L.UNIVERSITY_TEXT, "sg/ntu|uni", "Why included (max 400)", "A new Singapore reason.")
    result = run(root, workbook)
    assert result.problems == []
    assert json.loads((root / "data/universities/sg/ntu.json").read_text())["whyIncluded"] == "A new Singapore reason."
    assert [Path(w).name for w in result.written] == ["ntu.json"]


@pytest.mark.parametrize("sheet,key,header,value,message", [
    (L.FEES, "uk/manchester|0|cur|international|2027", "Annual tuition", -5, "negative"),
    (L.FEES, "uk/manchester|0|cur|international|2027", "Annual tuition", "lots", "must be a number"),
    (L.FEES, "uk/manchester|0|cur|international|2027", "Last verified", "21/10/2026", "must look like"),
    (L.FEES, "uk/manchester|0|cur|international|2027", "Source type", "blog", "official or secondary"),
    (L.FEES, "uk/manchester|0|cur|international|2027", "Source URL", "not a link", "http"),
    (L.PROGRAMME_TEXT, "uk/manchester|0|prog", "How this fee was worked out (shown, max 160, one sentence, no amounts)",
     "Uses the top of the range, £39,700 a year.", "amounts"),
    (L.PROGRAMME_TEXT, "uk/manchester|0|prog", "How this fee was worked out (shown, max 160, one sentence, no amounts)",
     "First sentence. Second sentence.", "one sentence"),
    (L.PROGRAMME_TEXT, "uk/manchester|0|prog", "Summary (shown, max 280)", "x" * 300, "limit is 280"),
    (L.UNIVERSITY_TEXT, "sg/ntu|uni", "Why included (max 400)", None, "cannot be empty"),
])
def test_bad_edits_are_reported_and_nothing_is_written(root, workbook, sheet, key, header, value, message):
    edit(workbook, sheet, key, header, value)
    before = snapshot(root)
    result = run(root, workbook)
    assert any(message in p for p in result.problems), result.problems
    assert snapshot(root) == before


def test_deleted_row_is_reported(root, workbook):
    wb = load_workbook(workbook)
    ws = wb[L.FEES]
    ws.delete_rows(find(ws, "uk/manchester|0|cur|international|2027"))
    wb.save(workbook)
    result = run(root, workbook)
    assert any("was deleted" in p for p in result.problems)


def test_row_with_wrong_identity_is_reported(root, workbook):
    edit(workbook, L.FEES, "uk/manchester|0|cur|international|2027", "Programme", "Something else")
    result = run(root, workbook)
    assert any("key belongs to" in p for p in result.problems)


def test_conflicting_source_values_on_one_programme(root, workbook):
    import json
    prog = json.loads((root / "data/universities/sg/ntu.json").read_text())["programmes"][0]
    year = prog["feeYear"]
    edit(workbook, L.FEES, f"sg/ntu|0|cur|citizen|{year}", "Source URL", "https://example.org/a")
    edit(workbook, L.FEES, f"sg/ntu|0|cur|pr|{year}", "Source URL", "https://example.org/b")
    result = run(root, workbook)
    assert any("different sourceUrl" in p for p in result.problems)
